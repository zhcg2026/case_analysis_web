# -*- coding: utf-8 -*-
"""考核月报 Word 生成：按"运行月报"红头公文版式输出 docx

版式取自 2026年7月月报 实测：
- 红头：单位名称 方正楷体_GBK 26pt 红色分散对齐；运 行 月 报 方正小标宋_GBK 65pt 红色居中
- 期号行：方正楷体_GB2312 16pt（期号 + 右对齐制表位日期）
- 一级标题 黑体16pt、（一）级 方正楷体_GB2312 16pt、1、级 楷体16pt，正文 仿宋_GB2312 16pt
- 表格：表头 548DD4 蓝底白色加粗 方正仿宋_GB2312 12pt，数据 仿宋_GB2312 12pt 居中
- 注：楷体 12pt 灰色 7E7E7E
"""
import logging
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm, Twips

try:
    from backend.assessment_report.report_data import work_note_lines, MUNICIPAL_ROWS
except ImportError:
    from .report_data import work_note_lines, MUNICIPAL_ROWS

try:
    from backend.assessment_routes import DISPATCH_TEAM_MAP, GARDEN_DISTRICT_MAP, PARK_MAP
except ImportError:
    from assessment_routes import DISPATCH_TEAM_MAP, GARDEN_DISTRICT_MAP, PARK_MAP

logger = logging.getLogger(__name__)

# 字体
F_TITLE_KAI = '方正楷体_GBK'       # 红头单位名称
F_TITLE_SONG = '方正小标宋_GBK'    # 红头 运行月报
F_KAI_GB = '方正楷体_GB2312'      # 期号行、（一）级标题、目录二级
F_FANGSONG = '仿宋_GB2312'        # 正文、表格数据
F_FANGSONG2 = '方正仿宋_GB2312'   # 图注、表头
F_HEI = '黑体'                    # 一级标题、目录一级
F_KAI = '楷体'                    # 1、级标题、注、工作动态

RED = RGBColor(0xFF, 0x00, 0x00)
GRAY = RGBColor(0x7E, 0x7E, 0x7E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = '548DD4'

CONTENT_W = 8500  # 版心宽度 twips（A4 - 左右3cm）


def _set_run(run, cn_font, size_pt, color=None, bold=False):
    run.font.name = cn_font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), cn_font)
    run.font.color.rgb = color if color is not None else BLACK


def _pf(p, align=None, indent_chars=0, size_pt=16, line_pt=None, before=0, after=0):
    """段落格式：首行缩进N字符、固定行距"""
    if align is not None:
        p.alignment = align
    pPr = p._p.get_or_add_pPr()
    if indent_chars:
        ind = pPr.find(qn('w:ind'))
        if ind is None:
            ind = OxmlElement('w:ind')
            pPr.append(ind)
        ind.set(qn('w:firstLineChars'), str(indent_chars * 100))
        ind.set(qn('w:firstLine'), str(int(indent_chars * size_pt * 20)))
    if line_pt:
        sp = pPr.find(qn('w:spacing'))
        if sp is None:
            sp = OxmlElement('w:spacing')
            pPr.append(sp)
        sp.set(qn('w:line'), str(int(line_pt * 20)))
        sp.set(qn('w:lineRule'), 'exact')
    if before or after:
        sp = pPr.find(qn('w:spacing'))
        if sp is None:
            sp = OxmlElement('w:spacing')
            pPr.append(sp)
        if before:
            sp.set(qn('w:before'), str(int(before * 20)))
        if after:
            sp.set(qn('w:after'), str(int(after * 20)))
    return p


def _add_para(doc, text, cn_font=F_FANGSONG, size_pt=16, align=None, indent_chars=2,
              line_pt=23, color=None, bold=False, before=0, after=0):
    p = doc.add_paragraph()
    _pf(p, align=align, indent_chars=indent_chars, size_pt=size_pt,
        line_pt=line_pt, before=before, after=after)
    run = p.add_run(text)
    _set_run(run, cn_font, size_pt, color=color, bold=bold)
    return p


def _set_cell_bg(cell, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), 'auto')
        borders.append(el)
    tblPr.append(borders)


def _set_row_cant_split(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement('w:cantSplit')
    trPr.append(cant)
    if header:
        th = OxmlElement('w:tblHeader')
        trPr.append(th)


def _set_cell_margins(table, top=15, left=15, bottom=15, right=15):
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for side, w in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(w))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblPr.append(mar)


def _norm_widths(widths, total=CONTENT_W):
    """保留参考文档各表的实际列宽网格。"""
    return list(widths)


def _add_table(doc, headers, rows, widths, header_size=12, data_size=12,
               header_font=F_FANGSONG2, data_font=F_FANGSONG):
    """通用表格：蓝底白字表头 + 仿宋数据行，全部居中"""
    widths = _norm_widths(widths)
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths)))
    tblW.set(qn('w:type'), 'dxa')
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'autofit')
    tblPr.append(layout)
    _set_table_borders(table)
    _set_cell_margins(table)

    # 表头
    hdr = table.rows[0]
    _set_row_cant_split(hdr, header=True)
    for j, h in enumerate(headers):
        cell = hdr.cells[j]
        cell.width = Twips(widths[j])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_bg(cell, HEADER_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        pPr = p._p.get_or_add_pPr()
        spacing = pPr.find(qn('w:spacing'))
        if spacing is None:
            spacing = OxmlElement('w:spacing')
            pPr.append(spacing)
        spacing.set(qn('w:line'), '240')
        spacing.set(qn('w:lineRule'), 'auto')
        parts = str(h).split('\n')
        for k, part in enumerate(parts):
            if k:
                p.add_run().add_break()
            run = p.add_run(part)
            _set_run(run, header_font, header_size, color=WHITE, bold=True)
    # 数据行（数字用 Times New Roman，度量稳定不因字体替换变化）
    for i, row_data in enumerate(rows):
        row = table.rows[i + 1]
        _set_row_cant_split(row)
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            cell.width = Twips(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            pPr = p._p.get_or_add_pPr()
            spacing = pPr.find(qn('w:spacing'))
            if spacing is None:
                spacing = OxmlElement('w:spacing')
                pPr.append(spacing)
            spacing.set(qn('w:line'), '240')
            spacing.set(qn('w:lineRule'), 'auto')
            run = p.add_run('' if val is None else str(val))
            _set_run(run, data_font, data_size)
            rFonts = run._element.get_or_add_rPr().find(qn('w:rFonts'))
            rFonts.set(qn('w:ascii'), 'Times New Roman')
            rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    return table


def _add_note(doc, lines):
    """表格下方灰色小字注"""
    for ln in lines:
        if ln:
            _add_para(doc, ln, cn_font=F_KAI, size_pt=12, indent_chars=2,
                      line_pt=16, color=GRAY)


def _add_caption(doc, text):
    _add_para(doc, text, cn_font=F_FANGSONG2, size_pt=16,
              align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0, line_pt=20)


def _add_red_rule(doc):
    """期号行下方的红色分隔线。"""
    p = doc.add_paragraph()
    _pf(p, line_pt=1, before=0, after=8, indent_chars=0)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'FF0000')
    borders.append(bottom)
    pPr.append(borders)
    return p


def _add_image(doc, img_path, caption, width_cm):
    p = doc.add_paragraph()
    _pf(p, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run()
    run.add_picture(img_path, width=Cm(width_cm))
    _add_caption(doc, caption)


def _fmt_score(v, dp=2):
    """分数显示：整数不带小数，其余固定小数位"""
    if v is None:
        return '—'
    v = float(v)
    if v == int(v):
        return str(int(v))
    return f'{v:.{dp}f}'


def _fmt_extra(v):
    """加减分项：去掉多余小数尾零（-0.1、-1、0）"""
    if v is None:
        return '—'
    v = float(v)
    if v == int(v):
        return str(int(v))
    return f'{v:.3f}'.rstrip('0').rstrip('.')


def _fmt_rate(n, d):
    if not d:
        return '-'
    return f'{n / d * 100:.2f}%'


# 各表列宽（twips，按参考月报保留各表独立网格）
W_DISPOSAL = [3450, 1288, 1204, 1284, 1269]
W_DISPATCH = [1489, 631, 775, 784, 796, 799, 600, 807, 681, 635, 605, 796]
W_SANITATION = [1256, 634, 600, 854, 766, 774, 813, 648, 704, 665, 717]
W_GARDEN = [1402, 618, 833, 853, 866, 842, 796, 704, 831, 615]
W_PARK = [1108, 635, 784, 831, 795, 810, 1090, 998, 831, 518]
W_MUNICIPAL = [1673, 996, 791, 832, 846, 755, 805, 641, 545]
W_PENDING = [2223, 873, 3460, 2167]
W_BACKLOG = [2188, 842, 5254]
W_PRAISE = [1412, 1754, 5118]


def _sorted_units(results, names):
    """按总分降序，空分/豁免排最后"""
    units = [n for n in names if n in results]
    units.sort(key=lambda n: (
        results[n].get('final_score') is None,
        -(results[n].get('final_score') or 0),
    ))
    return units


def _score_rows(data, unit_names, dp, extra_cols):
    """考核得分表通用行构造。extra_cols: callable(unit_name, r) -> list"""
    rows = []
    for name in unit_names:
        r = data['scores'].get(name, {})
        total = r.get('total', 0)
        rows.append([
            name, total or 0, r.get('closed', 0),
            _fmt_rate(r.get('overtime', 0), total),
            _fmt_rate(r.get('delayed', 0), total),
            _fmt_rate(r.get('rework', 0), total),
        ] + extra_cols(name, r))
    return rows


def _dispatch_extra(name, r):
    缓办率 = '0'
    return [
        缓办率,
        _fmt_score(r.get('system_score'), 3),
        _fmt_score(r.get('team_score'), 1),
        _fmt_score(r.get('street_score'), 1),
        _fmt_extra(r.get('extra_points')),
        _fmt_score(r.get('final_score'), 3),
    ]


def _sanitation_extra(name, r):
    return [
        '0',
        _fmt_score(r.get('system_score'), 2),
        _fmt_score(r.get('garbage_score'), 2),
        _fmt_score(r.get('center_score'), 1),
        _fmt_score(r.get('final_score'), 2),
    ]


def _garden_extra(data):
    def _cols(name, r):
        pending = data['pending_by_unit'].get(name, 0)
        return [
            _fmt_score(r.get('system_score'), 2),
            _fmt_score(r.get('center_score'), 1),
            _fmt_score(r.get('final_score'), 2),
            str(pending) if pending else '0',
        ]
    return _cols


def _municipal_rows(data):
    rows = []
    for key, display in MUNICIPAL_ROWS:
        r = data['scores'].get(key, {})
        if not r:
            continue
        total = r.get('total', 0)
        if r.get('score_mode') == 'close_rate':
            rate_cols = ['—', '—', '—', '—']
        else:
            rate_cols = [
                _fmt_rate(r.get('closed', 0), total),
                _fmt_rate(r.get('overtime', 0), total),
                _fmt_rate(r.get('delayed', 0), total),
                _fmt_rate(r.get('rework', 0), total),
            ]
        backlog = data['backlog_by_unit'].get(key, 0) or data['backlog_by_unit'].get(display, 0)
        rows.append([display, total or 0, r.get('closed', 0), *rate_cols,
                     _fmt_score(r.get('final_score'), 2),
                     str(backlog) if backlog else ''])
    return rows


def generate_word(data: dict, img1: str, img2: str, out_path: str) -> str:
    doc = Document()

    # 页面：首节对应参考文档的标题页，目录后用正文节承载报告内容。
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Twips(1440)
    sec.bottom_margin = Twips(1701)
    sec.left_margin = Twips(1701)
    sec.right_margin = Twips(1701)
    sec.header_distance = Twips(851)
    sec.footer_distance = Twips(1417)

    # ===== 红头 =====
    doc.add_paragraph()
    doc.add_paragraph()
    _add_para(doc, data['org_name'], cn_font=F_TITLE_KAI, size_pt=26,
              align=WD_ALIGN_PARAGRAPH.DISTRIBUTE, indent_chars=0, color=RED,
              line_pt=34, before=16, after=8)
    _add_para(doc, '运 行 月 报', cn_font=F_TITLE_SONG, size_pt=65,
              align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0, color=RED,
              line_pt=78, before=32, after=16)

    # 期号行：期号 + 右对齐制表位日期
    issue = data['issue']
    p = doc.add_paragraph()
    _pf(p, align=WD_ALIGN_PARAGRAPH.LEFT, size_pt=16, line_pt=28)
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:pos'), str(CONTENT_W))
    tabs.append(tab)
    pPr.append(tabs)
    run = p.add_run(f"〔{issue['year']}〕第{issue['no']}期（总第{issue['total']}期）")
    _set_run(run, F_KAI_GB, 16)
    run2 = p.add_run('\t' + issue['date_str'])
    _set_run(run2, F_KAI_GB, 16)
    _add_red_rule(doc)

    # ===== 目录块 =====
    _add_para(doc, '', size_pt=16, indent_chars=0, line_pt=14)
    toc_items = [
        ('综合运行情况', F_HEI, 1),
        ('（一）处置情况', F_KAI_GB, 2),
        ('（二）考核得分', F_KAI_GB, 2),
        ('挂帐、积压问题', F_HEI, 1),
        ('工作动态', F_HEI, 1),
        ('附件：数据分析', F_HEI, 1),
    ]
    for text, font, ind in toc_items:
        _add_para(doc, f'· {text}', cn_font=font, size_pt=16, indent_chars=ind, line_pt=28)

    body_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    body_sec.orientation = WD_ORIENT.PORTRAIT
    body_sec.page_width = Cm(21)
    body_sec.page_height = Cm(29.7)
    body_sec.top_margin = Twips(1383)
    body_sec.bottom_margin = Twips(1701)
    body_sec.left_margin = Twips(1701)
    body_sec.right_margin = Twips(1701)
    body_sec.header_distance = Twips(851)
    body_sec.footer_distance = Twips(1134)

    o = data['overview']
    received = o['received'] if o['received'] is not None else o['batch_total']
    received_total = (received or 0) + o['no_assess_total']

    # ===== 一、综合运行情况 =====
    _add_para(doc, '一、综合运行情况', cn_font=F_HEI, size_pt=16, indent_chars=2, line_pt=23)
    _add_para(doc, f"{data['period']}，{data['org_name']}共接收各来源渠道案件{received_total}件，"
                   f"其中考核案件{received}件，专项采集不考核{o['no_assess_total']}件。")
    _add_para(doc, f"1、系统考核（按处置结案）受理{received}件，应结案{o['batch_total']}件，"
                   f"结案{o['batch_closed']}件，结案率{o['close_rate']:.2f}%。")
    if img1:
        _add_image(doc, img1, '图1.问题来源立案占比', 15.48)
    if o['special_total']:
        detail = '、'.join(f'{n}{c}件' for n, c in o['special_breakdown'][:5])
        _add_para(doc, f"专项采集城市问题{o['special_total']}件，其中{detail}。")
    _add_para(doc, f"采集员自行处置{o['self_dispose']}件，单体垃圾扣分{o['garbage_cnt']}件。")
    if img2:
        _add_image(doc, img2, '图2.不考核案件类型及数量', 13.58)

    # （一）处置情况
    _add_para(doc, '（一）处置情况', cn_font=F_KAI_GB, size_pt=16, indent_chars=2, line_pt=23)
    disp = data['disposal']
    rows = [[r['name'], r['total'], r['closed'], f"{r['rate']:.2f}%", f"{r['share']:.2f}%"]
            for r in disp['rows']]
    t = disp['total']
    rows.append(['合  计', t['total'], t['closed'], f"{t['rate']:.2f}%", '100.00%'])
    _add_table(doc, ['处置部门', '应结案数', '结案数', '结案率', '占比'], rows, W_DISPOSAL)

    # （二）考核得分
    _add_para(doc, '（二）考核得分', cn_font=F_KAI_GB, size_pt=16, indent_chars=2, line_pt=23)

    results = data['scores']
    sys_formula = '其中系统考核得分=（按期结案率*100%+超期结案率*40%）*80%+（1-延期率）*10%+（1-返工率）*10%±加减分项'

    # 1、市容秩序
    _add_para(doc, '1、市容秩序', cn_font=F_KAI, size_pt=16, indent_chars=2, line_pt=22)
    teams = _sorted_units(results, list(DISPATCH_TEAM_MAP.values()))
    dispatch_rows = _score_rows(data, teams, 3, _dispatch_extra)
    _add_table(doc,
               ['执法\n分队', '应结\n案数', '结案数', '超期率', '延期率', '返工率', '缓办率',
                '系统\n分数', '队考\n核分', '街道\n办分', '扣分', '总分'],
               dispatch_rows, W_DISPATCH)
    notes = [f'注：1.计算公式：执法队各队考核得分=系统考核得分*70%+队考核得分*15%+街道办考核得分*15%；{sys_formula}。']
    if data.get('extra_note'):
        notes.append(f'2.{data["extra_note"]}')
    _add_note(doc, notes)

    # 2、环境卫生
    _add_para(doc, '2、环境卫生', cn_font=F_KAI, size_pt=16, indent_chars=2, line_pt=22)
    san_names = [n for n in results if '环卫' in n and n != '市容环卫中心']
    san_names = _sorted_units(results, san_names)
    san_rows = _score_rows(data, san_names, 2, _sanitation_extra)
    _add_table(doc,
               ['环卫\n片区', '应结\n案数', '结案数', '超期率', '延期率', '返工率', '缓办率',
                '系统\n分数', '单体\n垃圾', '中心\n分数', '总分'],
               san_rows, W_SANITATION)
    _add_note(doc, [
        f'注：计算公式：环卫考核得分=系统考核得分*30%+单体垃圾得分*30%+中心得分*40%；'
        f'其中单体垃圾得分=100-案件*0.01；{sys_formula}。'
    ])

    # 3、园林绿化
    _add_para(doc, '3、园林绿化', cn_font=F_KAI, size_pt=16, indent_chars=2, line_pt=22)
    garden_names = _sorted_units(results, list(GARDEN_DISTRICT_MAP.keys()))
    garden_rows = _score_rows(data, garden_names, 2, _garden_extra(data))
    _add_table(doc,
               ['园林\n片区', '应结\n案数', '结案数', '超期率', '延期率', '返工率',
                '系统\n分数', '中心\n分数', '总分', '挂账'],
               garden_rows, W_GARDEN)
    park_names = _sorted_units(results, list(PARK_MAP.keys()))
    park_rows = _score_rows(data, park_names, 2, _garden_extra(data))
    _add_table(doc,
               ['公园\n广场', '应结\n案数', '结案数', '超期率', '延期率', '返工率',
                '系统\n分数', '中心\n分数', '总分', '挂账'],
               park_rows, W_PARK)
    y, m = data['batch'][:4], int(data['batch'][4:6])
    _add_note(doc, [
        f'注：1.挂帐数为{m}月底前案件数；',
        f'2.计算公式：园林考核得分=系统考核得分*70%+中心考核得分*30%，{sys_formula}。'
    ])

    # 4、市政公用
    _add_para(doc, '4、市政公用', cn_font=F_KAI, size_pt=16, indent_chars=2, line_pt=22)
    muni_rows = _municipal_rows(data)
    _add_table(doc,
               ['市政\n考核', '应结\n案数', '结案数', '结案率', '超期率', '延期率', '返工率', '分数', '积压'],
               muni_rows, W_MUNICIPAL)
    _add_note(doc, [
        f'注：1、挂帐数为{m}月底前案件数；积压数为{y}年1月至今案件数。',
        f'计算公式：照明、排水系统考核得分=（按期结案率*100%+超期结案率*40%）*80%'
        f'+（1-延期率）*10%+（1-返工率）*10%±加减分项；应急、维护部暂按处置率得分。'
    ])

    # ===== 二、挂帐、积压问题 =====
    _add_para(doc, '二、挂帐、积压问题', cn_font=F_HEI, size_pt=16, indent_chars=2, line_pt=23)
    ledger = data['ledger']

    _add_para(doc, '1、部门挂帐总数', cn_font=F_KAI, size_pt=16, indent_chars=2, line_pt=22)
    if ledger['pending']:
        rows = [[it['unit_name'], it['piece_cnt'] if it['piece_cnt'] is not None else '',
                 it['content'] or '—',
                 ((it['reason'] or '') + (it['deadline'] or '')) or '—']
                for it in ledger['pending']]
        cnt_sum = sum(it['piece_cnt'] or 0 for it in ledger['pending'])
        rows.append(['合计', cnt_sum, '—', ''])
        _add_table(doc, ['部门', '数量', '主要内容', '理由、时间'], rows, W_PENDING,
                   header_size=14, data_font=F_FANGSONG2)
    else:
        _add_para(doc, '本月无。')

    _add_para(doc, f'2、1-{m}月积压数', cn_font=F_KAI, size_pt=16, indent_chars=2, line_pt=22)
    if ledger['backlog']:
        rows = [[it['dept_name'] or it['unit_name'], it['piece_cnt'] if it['piece_cnt'] is not None else '',
                 it['content'] or '—']
                for it in ledger['backlog']]
        cnt_sum = sum(it['piece_cnt'] or 0 for it in ledger['backlog'])
        rows.append(['合计', cnt_sum, '—'])
        _add_table(doc, ['部门', '数量', '主要内容'], rows, W_BACKLOG,
                   header_size=14, data_font=F_FANGSONG2)
    else:
        _add_para(doc, '本月无。')

    _add_para(doc, '3、表扬件', cn_font=F_KAI, size_pt=16, indent_chars=2, line_pt=22)
    if ledger['praise']:
        rows = [[it['source'] or '—', it['unit_name'] or '—', it['content'] or '—']
                for it in ledger['praise']]
        _add_table(doc, ['来源', '部门', '内容'], rows, W_PRAISE,
                   header_size=14, header_font=F_FANGSONG, data_font=F_FANGSONG2)
    else:
        _add_para(doc, '本月无。')

    # ===== 三、工作动态 =====
    _add_para(doc, '三、工作动态', cn_font=F_HEI, size_pt=16, indent_chars=2, line_pt=23)
    items = work_note_lines(data.get('work_note') or '')
    if items:
        for it in items:
            _add_para(doc, it, cn_font=F_KAI, size_pt=16, indent_chars=2, line_pt=25)
    else:
        _add_para(doc, '无。', cn_font=F_KAI, size_pt=16, indent_chars=2, line_pt=25)

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)
    logger.info(f'考核月报生成: {out_path}')
    return out_path
