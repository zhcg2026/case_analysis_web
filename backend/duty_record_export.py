# -*- coding: utf-8 -*-
"""值班记录 Word 导出 - 单班次(白班/夜班)与整月两种文档

版式沿用值班人员手工 txt 的格式：
  白班：`9月13日值班人员: 张三、李四` + 一、系统运行(3总数+7来源) + 二、关注问题
  夜班：`9月1日夜间值班：王康乐` + 一切正常 或 事件(1.12345（44）+ 来电人 + 时间线 + 处理结果)
"""
import io
import datetime
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

MIMETYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'

# 白班统计行（顺序与录入一致）
_STAT_LINES = [
    ('stat_reported', '上报'), ('stat_accepted', '受理'), ('stat_completed', '办结'),
]
_SRC_LINES = [
    ('src_collector', '采集员上报受理'),
    ('src_patrol', '重点领域日常巡查受理'),
    ('src_12345', '12345系统转办'),
    ('src_minhu', '民呼我应'),
    ('src_video', '视频监控'),
    ('src_ai', '智能分析'),
    ('src_public', '市民举报系统受理'),
]


def _new_doc(title):
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal.paragraph_format.space_after = Pt(4)
    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(16)
        run.font.name = 'Times New Roman'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    return doc


def _add(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    return p


def _d_cn(d):
    return f'{d.month}月{d.day}日'


def _shift_day_block(doc, d, rec, events):
    """白班段落：人员 + 一、系统运行 + 二、关注问题"""
    _add(doc, f'{_d_cn(d)}值班人员: {rec.members or "无"}', bold=True)
    stats = {f: getattr(rec, f) for f, _ in _STAT_LINES + _SRC_LINES}
    if any(stats[f] is not None for f, _ in _STAT_LINES):
        _add(doc, '一、系统运行', bold=True)
        _add(doc, f"上报{stats['stat_reported'] or 0}件，受理{stats['stat_accepted'] or 0}件，办结{stats['stat_completed'] or 0}件。")
        for f, label in _SRC_LINES:
            _add(doc, f'{label}:{stats[f] if stats[f] is not None else 0}')
    issues = [e for e in events if e['category'] == '关注问题']
    _add(doc, '二、关注问题', bold=True)
    if not issues:
        _add(doc, '无')
    for i, e in enumerate(issues, 1):
        _add(doc, f"{i}.{e['ticket_no'] or ''}".rstrip('.'))
        detail = '，'.join(x for x in (e['location'], e['description']) if x)
        if detail:
            _add(doc, detail + ('。' if not detail.endswith('。') else ''))
        for t in e['timeline']:
            _add(doc, f"{t.get('time', '')} {t.get('text', '')}".strip())
        if e['result']:
            _add(doc, f"案件详情：{e['result']}")
    if rec.note:
        _add(doc, f'备注：{rec.note}')


def _shift_night_block(doc, d, rec, events):
    """夜班段落：人员 + 一切正常 或 事件流水"""
    _add(doc, f'{_d_cn(d)}夜间值班：{rec.members or "无"}', bold=True)
    if rec.is_normal:
        _add(doc, '一切正常')
        if rec.note:
            _add(doc, f'备注：{rec.note}')
        return
    for i, e in enumerate(events, 1):
        if e['category'] == '12345':
            head = f'{i}.12345'
            if e['ticket_no']:
                head += f"（{e['ticket_no']}）"
        else:
            head = f"{i}.{e['category']}"
        _add(doc, head)
        caller = ' '.join(x for x in (e['caller_name'], e['caller_phone']) if x)
        if caller:
            _add(doc, f'来电人：{caller}')
        if e['location']:
            _add(doc, e['location'])
        if e['description']:
            text = e['description']
            _add(doc, text + ('。' if not text.endswith('。') else ''))
        for t in e['timeline']:
            _add(doc, f"{t.get('time', '')} {t.get('text', '')}".strip())
        if e['result']:
            _add(doc, f'处理结果：{e["result"]}')
    if rec.note:
        _add(doc, f'备注：{rec.note}')


def build_shift_doc(d, shift, rec, events):
    """单日单班次 docx 字节"""
    title = f'{d.year}年{_d_cn(d)} {shift}值班记录'
    doc = _new_doc(title)
    if shift == '白班':
        _shift_day_block(doc, d, rec, events)
    else:
        _shift_night_block(doc, d, rec, events)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_month_doc(year, month, day_records):
    """整月 docx 字节。day_records: 按日期升序的 [(date, shift, rec, events)]"""
    doc = _new_doc(f'{year}年{month}月值班记录')
    cur_date = None
    for d, shift, rec, events in day_records:
        if cur_date is not None and d != cur_date:
            doc.add_paragraph()  # 日期之间空一行
        cur_date = d
        if shift == '白班':
            _shift_day_block(doc, d, rec, events)
        else:
            _shift_night_block(doc, d, rec, events)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
